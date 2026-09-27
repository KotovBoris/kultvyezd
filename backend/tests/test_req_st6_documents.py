"""Требования ST-6 (docs/PRODUCT_GAPS.md №5, №6, №9; §6.4).

Пакет документов: маршрутный лист и уведомление в ГИБДД (ПП РФ №1527), согласие
на обработку ПДн (152-ФЗ), выбор состава приложений, отправка файла в MAX-чат,
номер билета.
"""
from __future__ import annotations

import asyncio
import io

from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import Session, select

from app.bot import BotService
from app.db import engine
from app.documents import ATTACHMENTS, normalize_attachments
from app.models import ConsentAudit

from conftest import make_excursion


def _docx_text(content: bytes) -> str:
    doc = Document(io.BytesIO(content))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _approved(client: TestClient, klass: dict, event_id: int | None = None, n: int = 2) -> dict:
    exc = make_excursion(client, klass, event_id)
    for s in klass["students"][:n]:
        client.post(f"/api/v1/excursions/{exc['id']}/consent",
                    json={"student_id": s["id"], "status": "APPROVED"})
    return exc


# --------------------------------------------------------- №5: маршрутный лист + ГИБДД
def test_ST6_attachments_catalog(client: TestClient) -> None:
    body = client.get("/api/v1/documents/attachments").json()
    keys = {a["key"] for a in body["attachments"]}
    assert {"students", "consents", "briefing", "route", "gibdd", "pdn"} <= keys


def test_ST6_docx_route_gibdd_present(client: TestClient, first_class, paid_event) -> None:
    exc = _approved(client, first_class, paid_event["id"])
    content = client.get(f"/api/v1/excursions/{exc['id']}/export-order",
                         params={"fmt": "docx", "attach": "students,route,gibdd,pdn"}).content
    text = _docx_text(content)
    assert "Маршрутный лист" in text
    assert "Госавтоинспекции" in text  # уведомление в ГИБДД
    assert "152-ФЗ" in text            # согласие на ПДн


def test_ST6_pdf_route_gibdd_ok(client: TestClient, first_class, paid_event) -> None:
    exc = _approved(client, first_class, paid_event["id"])
    r = client.get(f"/api/v1/excursions/{exc['id']}/export-order",
                   params={"fmt": "pdf", "attach": "route,gibdd,pdn"})
    assert r.status_code == 200 and r.content[:4] == b"%PDF"
    assert len(r.content) > 1000


# --------------------------------------------------------- выбор состава приложений
def test_ST6_attach_subset_honored(client: TestClient, first_class, paid_event) -> None:
    exc = _approved(client, first_class, paid_event["id"])
    body = client.post(f"/api/v1/excursions/{exc['id']}/export-order?attach=students,route").json()
    assert body["attachments"] == ["students", "route"]
    assert "attach=students,route" in body["docx_url"]

    content = client.get(f"/api/v1/excursions/{exc['id']}/export-order",
                         params={"fmt": "docx", "attach": "students,route"}).content
    text = _docx_text(content)
    assert "Список участников группы" in text
    assert "Маршрутный лист" in text
    assert "Лист целевого инструктажа" not in text  # приложение снято галочкой


def test_ST6_attach_unknown_keys_ignored() -> None:
    assert normalize_attachments(["route", "nonsense"]) == ["route"]
    assert normalize_attachments(None) == list(ATTACHMENTS)  # пусто → все приложения


def test_ST6_gibdd_can_be_disabled(client: TestClient, first_class, paid_event) -> None:
    exc = _approved(client, first_class, paid_event["id"])
    content = client.get(f"/api/v1/excursions/{exc['id']}/export-order",
                         params={"fmt": "docx", "attach": "students"}).content
    text = _docx_text(content)
    assert "Уведомление в подразделение Госавтоинспекции" not in text
    assert "Маршрутный лист" not in text


# --------------------------------------------------------- №6: согласие на ПДн (152-ФЗ)
def test_ST6_pdn_consent_saved(client: TestClient, first_class, paid_event) -> None:
    exc = make_excursion(client, first_class, paid_event["id"])
    sid = first_class["students"][0]["id"]
    client.post(f"/api/v1/excursions/{exc['id']}/consent",
                json={"student_id": sid, "status": "APPROVED", "pdn_consent": True,
                      "parent_phone": "+79001110011"})
    d = client.get(f"/api/v1/excursions/{exc['id']}/dashboard").json()
    row = next(p for p in d["participants"] if p["student_id"] == sid)
    assert row["pdn_consent_at"] and row["pdn_consent_by"] == "+79001110011"


def test_ST6_pdn_consent_absent_without_flag(client: TestClient, first_class, paid_event) -> None:
    exc = make_excursion(client, first_class, paid_event["id"])
    sid = first_class["students"][1]["id"]
    client.post(f"/api/v1/excursions/{exc['id']}/consent",
                json={"student_id": sid, "status": "APPROVED"})
    d = client.get(f"/api/v1/excursions/{exc['id']}/dashboard").json()
    row = next(p for p in d["participants"] if p["student_id"] == sid)
    assert row["pdn_consent_at"] is None


def test_ST6_pdn_consent_audited(client: TestClient, first_class, paid_event) -> None:
    exc = make_excursion(client, first_class, paid_event["id"])
    sid = first_class["students"][0]["id"]
    client.post(f"/api/v1/excursions/{exc['id']}/consent",
                json={"student_id": sid, "status": "APPROVED", "pdn_consent": True})
    with Session(engine) as s:
        rows = s.exec(select(ConsentAudit).where(ConsentAudit.excursion_id == exc["id"],
                                                 ConsentAudit.student_id == sid)).all()
    assert any(r.action == "PDN_CONSENT" for r in rows)


def test_ST6_reset_clears_pdn(client: TestClient) -> None:
    """reset-demo (первый, сидированный выезд) сбрасывает и отметку согласия на ПДн."""
    exc = client.get("/api/v1/excursions").json()[0]
    d = client.get(f"/api/v1/excursions/{exc['id']}/dashboard").json()
    sid = d["participants"][0]["student_id"]
    client.post(f"/api/v1/excursions/{exc['id']}/consent",
                json={"student_id": sid, "status": "APPROVED", "pdn_consent": True})
    mid = client.get(f"/api/v1/excursions/{exc['id']}/dashboard").json()
    assert next(p for p in mid["participants"] if p["student_id"] == sid)["pdn_consent_at"]
    client.post("/api/v1/admin/reset-demo")
    after = client.get(f"/api/v1/excursions/{exc['id']}/dashboard").json()
    assert next(p for p in after["participants"] if p["student_id"] == sid)["pdn_consent_at"] is None


# --------------------------------------------------------- №9: номер билета
def test_ST6_ticket_number_saved(client: TestClient, first_class, paid_event) -> None:
    exc = make_excursion(client, first_class, paid_event["id"])
    sid = first_class["students"][0]["id"]
    client.post(f"/api/v1/excursions/{exc['id']}/consent", json={"student_id": sid, "status": "APPROVED"})
    client.post(f"/api/v1/excursions/{exc['id']}/ticket-confirm",
                json={"student_id": sid, "ticket_number": "KZ-777"})
    d = client.get(f"/api/v1/excursions/{exc['id']}/dashboard").json()
    row = next(p for p in d["participants"] if p["student_id"] == sid)
    assert row["ticket_number"] == "KZ-777"


# --------------------------------------------------------- §6.4: отправка файла в MAX
def test_ST6_send_order_endpoint_no_bot(client: TestClient, first_class, paid_event) -> None:
    """Бот выключен (нет токена): endpoint не падает и честно сообщает, что не доставлено."""
    exc = _approved(client, first_class, paid_event["id"])
    r = client.post(f"/api/v1/excursions/{exc['id']}/send-order", params={"chat_id": -700})
    assert r.status_code == 200
    body = r.json()
    assert body["delivered"] is False and body["download_url"].startswith("http")
    assert body["bot_enabled"] is False


def test_ST6_send_order_requires_target(client: TestClient, first_class) -> None:
    exc = make_excursion(client, first_class)
    assert client.post(f"/api/v1/excursions/{exc['id']}/send-order").status_code == 422


def test_ST6_send_order_unknown_404(client: TestClient) -> None:
    assert client.post("/api/v1/excursions/999999/send-order", params={"chat_id": 1}).status_code == 404


def test_ST6_send_order_file_uploads_and_attaches(client: TestClient, first_class, paid_event) -> None:
    """С фейковым MAX: файл загружается (uploads) и уходит вложением type=file."""
    exc = _approved(client, first_class, paid_event["id"])
    from app.db import engine as _engine
    from app.models import Excursion

    with Session(_engine) as s:
        obj = s.get(Excursion, exc["id"])
        fake = _FakeUploadClient()
        svc = BotService(client=fake)
        delivered = asyncio.new_event_loop().run_until_complete(
            svc.send_order_file(obj, fmt="docx", chat_id=-700)
        )
    assert delivered is True
    assert fake.uploaded and fake.uploaded[0][1].endswith(".docx")
    assert fake.sent[-1]["attachments"][0]["type"] == "file"
    assert fake.sent[-1]["attachments"][0]["payload"]["token"] == "TOKEN123"


# --------------------------------------------------------------- helpers
class _FakeUploadClient:
    def __init__(self) -> None:
        self.uploaded: list[tuple[str, str]] = []
        self.sent: list[dict] = []

    @property
    def enabled(self) -> bool:
        return True

    async def upload_file(self, content: bytes, filename: str, upload_type: str = "file") -> str:
        self.uploaded.append((upload_type, filename))
        return "TOKEN123"

    async def send_document(self, *, filename: str, token: str, caption: str = "",
                            chat_id: int | None = None, user_id: int | None = None) -> dict:
        self.sent.append({
            "chat_id": chat_id, "user_id": user_id, "text": caption,
            "attachments": [{"type": "file", "payload": {"token": token, "filename": filename}}],
        })
        return {"ok": True}
