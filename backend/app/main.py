from __future__ import annotations

import asyncio
import csv
import io
import logging
import secrets
from contextlib import asynccontextmanager
from datetime import date, datetime, time, timezone

from fastapi import Depends, FastAPI, File, HTTPException, Query, Request, Response, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import JSONResponse
from openpyxl import load_workbook
from sqlmodel import Session, select

from . import __version__, demo, services
from .auth import Caller, caller_from_request, get_caller, remember_profile
from .bot import BotService, class_chat_id
from .catalog_seed import SOURCE, refresh_catalog_dates, seed_all
from .config import get_settings
from .db import Session, engine, get_session, init_db
from .documents import build_order_docx, build_order_pdf
from .max_client import BOT_COMMANDS
from .max_validate import is_fresh, validate_webapp_data, verify_phone
from .models import (
    AccessRole,
    BotLinkCode,
    ClassAccess,
    CultureEvent,
    DocumentArtifact,
    Excursion,
    ExcursionParticipant,
    ExcursionStatus,
    ParentContact,
    School,
    SchoolAccess,
    SchoolClass,
    Student,
    TicketStatus,
    UserProfile,
)
from .schemas import (
    ClassCreate,
    ClassOut,
    ClassUpdate,
    ConsentRequest,
    ConsentResetRequest,
    ConsentResponse,
    CultureEventOut,
    DashboardOut,
    DemoOut,
    DocumentOut,
    ExcursionCreate,
    ExcursionOut,
    ExcursionStats,
    ImportResult,
    InviteOut,
    AttendanceRequest,
    MilestoneOut,
    MilestoneRequest,
    PairInfoOut,
    PairRequest,
    SecondParentInviteOut,
    SecondParentInviteRequest,
    JoinInfoOut,
    JoinOut,
    JoinRequest,
    MeOut,
    RoleRequest,
    SendDocumentsOut,
    LinkCodeOut,
    NotificationsRequest,
    ParentClaimRequest,
    ParentLinkAction,
    ParentOut,
    ParentSearchRow,
    ParticipantRow,
    RemindResponse,
    SchoolCreate,
    SchoolOut,
    ShareOut,
    ShareRequest,
    StudentCreate,
    StudentOut,
    TicketConfirmRequest,
)

log = logging.getLogger("classgo")
settings = get_settings()

_poll_task: asyncio.Task | None = None
_reminder_task: asyncio.Task | None = None
_demo_task: asyncio.Task | None = None


def _require_valid_parent(init_data: str | None) -> None:
    if not settings.MAX_VALIDATE_INIT_DATA:
        return
    valid, _ = validate_webapp_data(init_data or "", settings.MAX_BOT_TOKEN)
    if not valid:
        raise HTTPException(status_code=401, detail="Не удалось проверить вход. Закройте и снова откройте ClassGo в MAX.")


def _naive_utc(value: datetime | None) -> datetime | None:
    """Дедлайн из браузера приходит с часовым поясом — храним в UTC без зоны, как datetime.utcnow()."""
    if value is None or value.tzinfo is None:
        return value
    return value.astimezone(timezone.utc).replace(tzinfo=None)


def _time_from_str(value: str | None) -> time | None:
    if not value:
        return None
    try:
        hh, mm = value.split(":")[:2]
        return time(int(hh), int(mm))
    except Exception:  # noqa: BLE001
        return None


def _excursion_out(exc: Excursion, session: Session | None = None, with_stats: bool = False) -> ExcursionOut:
    klass = session.get(SchoolClass, exc.class_id) if session is not None else None
    stats = None
    if with_stats and session is not None:
        data = services.excursion_summary(session, exc)
        stats = ExcursionStats(**{k: data[k] for k in ExcursionStats.model_fields})
    return ExcursionOut(
        id=exc.id, class_id=exc.class_id, title=exc.title, location_name=exc.location_name,
        address=exc.address, event_date=exc.event_date,
        gathering_time=exc.gathering_time.strftime("%H:%M") if exc.gathering_time else None,
        return_time=exc.return_time.strftime("%H:%M") if exc.return_time else None,
        deadline=exc.deadline, ticket_price=exc.ticket_price, ticket_sale_url=exc.ticket_sale_url,
        is_pushkin_card=exc.is_pushkin_card, status=exc.status, school_name=exc.school_name,
        responsible_teacher=exc.responsible_teacher, meeting_point=exc.meeting_point or "",
        escorts=exc.escorts or "", transport=exc.transport or "", departed_at=exc.departed_at,
        returned_at=exc.returned_at,
        cancelled=exc.cancelled,
        class_title=klass.title if klass else "", is_demo=bool(klass and klass.is_demo), stats=stats,
    )


def _user_class_ids(session: Session, user_id: int | None) -> list[int]:
    """Классы, которые пользователь ведёт или к которым ему открыли доступ."""
    if not user_id:
        return []
    owned = {k.id for k in session.exec(select(SchoolClass).where(SchoolClass.owner_user_id == user_id)).all()}
    shared = {a.class_id for a in session.exec(select(ClassAccess).where(ClassAccess.user_id == user_id)).all()}
    return sorted(owned | shared)


def _parent_out(p: ParentContact) -> ParentOut:
    return ParentOut(
        id=p.id, full_name=p.full_name, phone_number=p.phone_number, role=p.role,
        confirmed=p.confirmed, bot_activated=bool(p.max_user_id or p.is_virtual),
        notifications_enabled=p.notifications_enabled,
    )


def _student_out(session: Session, s: Student) -> StudentOut:
    parents = session.exec(select(ParentContact).where(ParentContact.student_id == s.id)).all()
    return StudentOut(
        id=s.id, full_name=s.full_name, birth_date=s.birth_date,
        parent_phones=[p.phone_number for p in parents],
        parents=[_parent_out(p) for p in parents],
    )


def _can_edit_class(session: Session, klass: SchoolClass, user_id: int | None) -> bool:
    if klass.owner_user_id is None:
        return True
    if user_id is None:
        return False
    if klass.owner_user_id == user_id:
        return True
    access = session.exec(
        select(ClassAccess).where(ClassAccess.class_id == klass.id, ClassAccess.user_id == user_id)
    ).first()
    return bool(access and access.role in (AccessRole.OWNER, AccessRole.EDIT))


def _can_edit_school(session: Session, school: School, user_id: int | None) -> bool:
    if school.owner_user_id is None:
        return True
    if user_id is None:
        return False
    if school.owner_user_id == user_id:
        return True
    access = session.exec(
        select(SchoolAccess).where(SchoolAccess.school_id == school.id, SchoolAccess.user_id == user_id)
    ).first()
    return bool(access and access.role in (AccessRole.OWNER, AccessRole.EDIT))


def _class_out(session: Session, k: SchoolClass, user_id: int | None = None) -> ClassOut:
    students = session.exec(select(Student).where(Student.class_id == k.id).order_by(Student.full_name)).all()
    student_rows = [_student_out(session, s) for s in students]
    connected = 0
    for s in students:
        parents = session.exec(select(ParentContact).where(ParentContact.student_id == s.id)).all()
        if any(p.max_user_id or p.is_virtual for p in parents):
            connected += 1
    return ClassOut(
        id=k.id, title=k.title, school_id=k.school_id, school_number=k.school_number,
        school_name=k.school_name, teacher_name=k.teacher_name, teacher_phone=k.teacher_phone or "",
        chat_id=k.chat_id, owner_user_id=k.owner_user_id, can_edit=_can_edit_class(session, k, user_id),
        is_demo=k.is_demo, invite_code=k.invite_code, parents_connected=connected,
        students=student_rows,
    )


async def _poll_loop() -> None:
    bot = BotService()
    me = await bot.client.get_me()
    log.info("MAX polling запущен, бот: %s", (me or {}).get("name"))
    registered = await bot.client.set_commands(BOT_COMMANDS)
    if registered is not None:
        log.info("Команды бота зарегистрированы в MAX: %s", ", ".join(c["name"] for c in BOT_COMMANDS))
    marker: int | None = None
    while True:
        try:
            data = await bot.client.get_updates(marker=marker, timeout=settings.POLL_TIMEOUT,
                                                types=["message_created", "message_callback", "bot_started"])
            marker = data.get("marker", marker)
            for update in data.get("updates") or []:
                await bot.handle_update(update)
        except asyncio.CancelledError:
            log.info("MAX polling остановлен")
            raise
        except Exception as exc:  # noqa: BLE001
            log.warning("Сбой polling: %s", exc)
            await asyncio.sleep(3)


async def _demo_loop() -> None:
    """Ответы виртуальных родителей демо-классов — каждые пару секунд."""
    bot = BotService()
    while True:
        try:
            with Session(engine) as session:
                touched = demo.process_due_answers(session)
            for excursion_id in touched:
                if settings.bot_enabled:
                    await bot.notify_teacher_if_complete(excursion_id)
            await asyncio.sleep(2)
        except asyncio.CancelledError:
            raise
        except Exception as exc:  # noqa: BLE001
            log.warning("Сбой цикла демо-ответов: %s", exc)
            await asyncio.sleep(5)


async def _reminder_loop() -> None:
    bot = BotService()
    while True:
        try:
            with Session(engine) as session:
                due = services.due_auto_reminders(session)
                for target in due:
                    sent = await bot.send_auto_reminder(target)
                    if sent:
                        services.mark_reminder_sent(
                            session, target["excursion"].id, target["student_id"], target["kind"]
                        )
            await asyncio.sleep(settings.REMINDER_INTERVAL_SECONDS)
        except asyncio.CancelledError:
            raise
        except Exception as exc:  # noqa: BLE001
            log.warning("Сбой цикла напоминаний: %s", exc)
            await asyncio.sleep(30)


@asynccontextmanager
async def lifespan(app: FastAPI):
    global _poll_task, _reminder_task, _demo_task
    logging.basicConfig(level=getattr(logging, settings.LOG_LEVEL, logging.INFO))
    init_db()
    if settings.AUTO_SEED:
        with Session(engine) as session:
            seed_all(session)
            refresh_catalog_dates(session)
    if settings.bot_enabled and settings.MAX_BOT_MODE == "polling":
        _poll_task = asyncio.create_task(_poll_loop())
    if settings.bot_enabled:
        _reminder_task = asyncio.create_task(_reminder_loop())
    _demo_task = asyncio.create_task(_demo_loop())
    yield
    for task in (_poll_task, _reminder_task, _demo_task):
        if task:
            task.cancel()
            try:
                await task
            except asyncio.CancelledError:
                pass


app = FastAPI(
    title="ClassGo API",
    description=(
        "Сервис организации школьных культурных выездов. "
        "Данные каталога — модельные (см. поле source), внешняя интеграция с PRO.Культура.РФ заглушена."
    ),
    version=__version__,
    lifespan=lifespan,
)
app.add_middleware(
    CORSMiddleware, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"],
)


@app.get("/healthz", tags=["service"])
def healthz() -> dict:
    return {"status": "ok", "version": __version__, "bot_mode": settings.MAX_BOT_MODE,
            "bot_enabled": settings.bot_enabled, "time": datetime.utcnow().isoformat()}


@app.get("/api/v1/meta", tags=["service"])
def meta() -> dict:
    return {
        "solution": "ClassGo",
        "version": __version__,
        "dataset": {
            "catalog_source": SOURCE,
            "provenance": "модельные данные; реальная интеграция требует доступа к API PRO.Культура.РФ",
            "is_mock": True,
        },
        "roles": ["teacher", "parent"],
        "max": {"api_base": settings.MAX_API_BASE, "bot_username": settings.MAX_BOT_USERNAME},
    }


@app.get("/api/v1/culture-events", response_model=list[CultureEventOut], tags=["catalog"])
def list_culture_events(
    city: str | None = Query(default=None),
    age: str | None = Query(default=None, description="Возрастной ценз: 0+, 6+, 12+, 16+"),
    pushkin: bool | None = Query(default=None),
    free: bool | None = Query(default=None),
    session: Session = Depends(get_session),
) -> list[CultureEvent]:
    stmt = select(CultureEvent)
    if city:
        stmt = stmt.where(CultureEvent.city == city)
    if age:
        stmt = stmt.where(CultureEvent.age_rating == age)
    if pushkin is not None:
        stmt = stmt.where(CultureEvent.pushkin_eligible == pushkin)  # noqa: E712
    if free is not None:
        stmt = stmt.where(CultureEvent.is_free == free)  # noqa: E712
    return list(session.exec(stmt).all())


@app.get("/api/v1/schools", response_model=list[SchoolOut], tags=["schools"])
def list_schools(query: str | None = Query(default=None), session: Session = Depends(get_session)) -> list[SchoolOut]:
    stmt = select(School)
    schools = session.exec(stmt).all()
    if query:
        q = query.strip().lower()
        schools = [s for s in schools if q in s.name.lower() or q in s.city.lower() or q == s.number]
    out = []
    for s in schools:
        count = len(session.exec(select(SchoolClass).where(SchoolClass.school_id == s.id)).all())
        out.append(SchoolOut(id=s.id, name=s.name, city=s.city, number=s.number,
                             owner_user_id=s.owner_user_id, classes_count=count))
    return out


@app.post("/api/v1/schools", response_model=SchoolOut, status_code=201, tags=["schools"])
def create_school(payload: SchoolCreate, session: Session = Depends(get_session)) -> SchoolOut:
    name = payload.name.strip()
    if not name:
        raise HTTPException(status_code=422, detail="Укажите название школы")
    existing = session.exec(select(School).where(School.name == name, School.city == payload.city)).first()
    if existing:
        return SchoolOut(id=existing.id, name=existing.name, city=existing.city, number=existing.number,
                         owner_user_id=existing.owner_user_id)
    school = School(name=name, city=payload.city, number=payload.number, owner_user_id=payload.user_id)
    session.add(school)
    session.commit()
    session.refresh(school)
    if payload.user_id:
        session.add(SchoolAccess(school_id=school.id, user_id=payload.user_id, role=AccessRole.OWNER))
        session.commit()
    return SchoolOut(id=school.id, name=school.name, city=school.city, number=school.number,
                     owner_user_id=school.owner_user_id)


@app.post("/api/v1/schools/{school_id}/share", response_model=ShareOut, tags=["schools"])
def share_school(school_id: int, payload: ShareRequest,
                 caller: Caller = Depends(get_caller),
                 session: Session = Depends(get_session)) -> ShareOut:
    school = session.get(School, school_id)
    if not school:
        raise HTTPException(status_code=404, detail="Школа не найдена")
    if not _can_edit_school(session, school, caller.user_id):
        raise HTTPException(status_code=403, detail="Нет доступа к этой школе")
    existing = session.exec(
        select(SchoolAccess).where(SchoolAccess.school_id == school_id, SchoolAccess.user_id == payload.user_id)
    ).first()
    if existing:
        existing.role = AccessRole(payload.role)
        session.add(existing)
    else:
        session.add(SchoolAccess(school_id=school_id, user_id=payload.user_id, role=AccessRole(payload.role)))
    session.commit()
    return ShareOut(user_id=payload.user_id, role=AccessRole(payload.role))


def _parse_rows(content: bytes, filename: str) -> list[dict]:
    rows: list[dict] = []
    if filename.lower().endswith(".xlsx") or filename.lower().endswith(".xlsm"):
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        headers = [str(c.value).strip().lower() if c.value else "" for c in next(ws.iter_rows(min_row=1, max_row=1))]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if not any(r):
                continue
            rows.append(dict(zip(headers, [("" if v is None else str(v).strip()) for v in r])))
    else:
        text = content.decode("utf-8-sig", errors="replace")
        # Русский Excel сохраняет CSV с «;» — определяем разделитель по первой строке.
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        for r in reader:
            rows.append({(k or "").strip().lower(): (v or "").strip() for k, v in r.items()})
    return rows


def _pick(row: dict, *keys: str) -> str:
    for k in keys:
        if k in row and row[k]:
            return str(row[k]).strip()
    return ""


def _apply_school(session: Session, klass: SchoolClass, school_id: int | None) -> None:
    if school_id is None:
        return
    school = session.get(School, school_id)
    if not school:
        raise HTTPException(status_code=404, detail="Школа не найдена")
    klass.school_id = school.id
    klass.school_name = school.name
    klass.school_number = school.number


@app.post("/api/v1/classes/import", response_model=ImportResult, tags=["classes"])
async def import_class(
    file: UploadFile = File(...),
    grade: str = Query(default="8"),
    letter: str = Query(default="А"),
    school_id: int | None = Query(default=None),
    school_number: str = Query(default=""),
    school_name: str = Query(default=""),
    teacher_name: str = Query(default=""),
    teacher_phone: str = Query(default=""),
    caller: Caller = Depends(get_caller),
    session: Session = Depends(get_session),
) -> ImportResult:
    content = await file.read()
    if not content:
        raise HTTPException(status_code=400, detail="Файл пустой")
    try:
        rows = _parse_rows(content, file.filename or "upload.csv")
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Не удалось разобрать файл: {exc}") from exc
    if not rows:
        raise HTTPException(status_code=400, detail="В файле нет строк с учениками")

    klass = SchoolClass(grade=grade, letter=letter, school_number=school_number,
                        school_name=school_name, teacher_name=teacher_name, teacher_phone=teacher_phone.strip(),
                        owner_user_id=caller.user_id)
    _apply_school(session, klass, school_id)
    session.add(klass)
    session.commit()
    session.refresh(klass)
    if caller.user_id:
        session.add(ClassAccess(class_id=klass.id, user_id=caller.user_id, role=AccessRole.OWNER))
        session.commit()

    students_created = parents_created = 0
    skipped: list[str] = []
    for i, row in enumerate(rows, start=2):
        full_name = _pick(row, "фио", "фио ребенка", "full_name", "ученик", "имя")
        if not full_name:
            skipped.append(f"строка {i}: нет ФИО")
            continue
        birth_raw = _pick(row, "дата рождения", "дата_рождения", "birth_date", "др")
        birth = None
        for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                birth = datetime.strptime(birth_raw, fmt).date()
                break
            except Exception:  # noqa: BLE001
                continue
        student = Student(class_id=klass.id, full_name=full_name, birth_date=birth)
        session.add(student)
        session.commit()
        session.refresh(student)
        students_created += 1
        phone = _pick(row, "телефон родителя", "телефон", "phone", "телефон мамы", "телефон папы")
        if phone:
            session.add(ParentContact(student_id=student.id, phone_number=phone,
                                      full_name=_pick(row, "фио родителя", "родитель", "parent")
                                      or f"Родитель {full_name.split()[0]}",
                                      role="Законный представитель"))
            parents_created += 1
    session.commit()
    return ImportResult(class_id=klass.id, title=klass.title, students_created=students_created,
                        parents_created=parents_created, skipped_rows=skipped[:20])


@app.post("/api/v1/classes", response_model=ClassOut, status_code=201, tags=["classes"])
def create_class(payload: ClassCreate, caller: Caller = Depends(get_caller),
                 session: Session = Depends(get_session)) -> ClassOut:
    if caller.verified or not payload.user_id:
        payload.user_id = caller.user_id
    if not payload.grade.strip():
        raise HTTPException(status_code=422, detail="Укажите номер класса")
    klass = SchoolClass(
        grade=payload.grade.strip(), letter=payload.letter.strip(),
        school_number=payload.school_number, school_name=payload.school_name,
        teacher_name=payload.teacher_name, teacher_phone=payload.teacher_phone.strip(), chat_id=payload.chat_id,
        owner_user_id=payload.user_id,
    )
    _apply_school(session, klass, payload.school_id)
    session.add(klass)
    session.commit()
    session.refresh(klass)
    if payload.user_id:
        session.add(ClassAccess(class_id=klass.id, user_id=payload.user_id, role=AccessRole.OWNER))
    for item in payload.students:
        student = Student(class_id=klass.id, full_name=item.full_name, birth_date=item.birth_date)
        session.add(student)
        session.commit()
        session.refresh(student)
        if item.parent_phone:
            session.add(ParentContact(
                student_id=student.id, phone_number=item.parent_phone,
                full_name=item.parent_name or f"Родитель {item.full_name.split()[0]}",
                role=item.parent_role,
            ))
    session.commit()
    return _class_out(session, klass, payload.user_id)


@app.get("/api/v1/classes", response_model=list[ClassOut], tags=["classes"])
def list_classes(mine: bool = Query(default=False, description="Только классы текущего пользователя"),
                 caller: Caller = Depends(get_caller),
                 session: Session = Depends(get_session)) -> list[ClassOut]:
    # mine=false — общий журнал (веб-проверка, совместимость); mine=true — классы учителя.
    if mine:
        ids = _user_class_ids(session, caller.user_id)
        classes = [session.get(SchoolClass, i) for i in ids]
    else:
        classes = session.exec(select(SchoolClass)).all()
    return [_class_out(session, k, caller.user_id) for k in classes if k]


@app.get("/api/v1/classes/{class_id}/roster", response_model=ClassOut, tags=["classes"])
def class_roster(class_id: int, caller: Caller = Depends(get_caller),
                 session: Session = Depends(get_session)) -> ClassOut:
    k = session.get(SchoolClass, class_id)
    if not k:
        raise HTTPException(status_code=404, detail="Класс не найден")
    return _class_out(session, k, caller.user_id)


@app.patch("/api/v1/classes/{class_id}", response_model=ClassOut, tags=["classes"])
def update_class(class_id: int, payload: ClassUpdate, caller: Caller = Depends(get_caller),
                 session: Session = Depends(get_session)) -> ClassOut:
    if caller.verified or not payload.user_id:
        payload.user_id = caller.user_id
    k = session.get(SchoolClass, class_id)
    if not k:
        raise HTTPException(status_code=404, detail="Класс не найден")
    if not _can_edit_class(session, k, payload.user_id):
        raise HTTPException(status_code=403, detail="Нет доступа к этому классу")
    if payload.grade is not None:
        k.grade = payload.grade.strip()
    if payload.letter is not None:
        k.letter = payload.letter.strip()
    if payload.teacher_name is not None:
        k.teacher_name = payload.teacher_name
    if payload.chat_id is not None:
        k.chat_id = payload.chat_id or None
    _apply_school(session, k, payload.school_id)
    session.add(k)
    session.commit()
    session.refresh(k)
    return _class_out(session, k, payload.user_id)


@app.post("/api/v1/classes/{class_id}/students", response_model=StudentOut, status_code=201, tags=["classes"])
def add_student(class_id: int, payload: StudentCreate,
                caller: Caller = Depends(get_caller),
                session: Session = Depends(get_session)) -> StudentOut:
    k = session.get(SchoolClass, class_id)
    if not k:
        raise HTTPException(status_code=404, detail="Класс не найден")
    if not _can_edit_class(session, k, caller.user_id):
        raise HTTPException(status_code=403, detail="Нет доступа к этому классу")
    if not payload.full_name.strip():
        raise HTTPException(status_code=422, detail="Укажите ФИО ученика")
    student = Student(class_id=class_id, full_name=payload.full_name.strip(), birth_date=payload.birth_date)
    session.add(student)
    session.commit()
    session.refresh(student)
    if payload.parent_phone:
        session.add(ParentContact(
            student_id=student.id, phone_number=payload.parent_phone,
            full_name=payload.parent_name or f"Родитель {student.full_name.split()[0]}",
            role=payload.parent_role,
        ))
        session.commit()
    for exc in session.exec(select(Excursion).where(Excursion.class_id == class_id)).all():
        if not services.get_participant(session, exc.id, student.id):
            session.add(ExcursionParticipant(
                excursion_id=exc.id, student_id=student.id,
                consent_status="PENDING",
                ticket_status=TicketStatus.NOT_REQUIRED if (exc.ticket_price or 0) == 0 else TicketStatus.WAITING_PAYMENT,
            ))
    session.commit()
    return _student_out(session, student)


@app.delete("/api/v1/classes/{class_id}/students/{student_id}", tags=["classes"])
def remove_student(class_id: int, student_id: int,
                   caller: Caller = Depends(get_caller),
                   session: Session = Depends(get_session)) -> dict:
    k = session.get(SchoolClass, class_id)
    if not k:
        raise HTTPException(status_code=404, detail="Класс не найден")
    if not _can_edit_class(session, k, caller.user_id):
        raise HTTPException(status_code=403, detail="Нет доступа к этому классу")
    student = session.get(Student, student_id)
    if not student or student.class_id != class_id:
        raise HTTPException(status_code=404, detail="Ученик не найден в этом классе")
    for p in session.exec(select(ExcursionParticipant).where(ExcursionParticipant.student_id == student_id)).all():
        session.delete(p)
    for pc in session.exec(select(ParentContact).where(ParentContact.student_id == student_id)).all():
        session.delete(pc)
    session.delete(student)
    session.commit()
    return {"ok": True, "student_id": student_id}


@app.post("/api/v1/classes/{class_id}/share", response_model=ShareOut, tags=["classes"])
def share_class(class_id: int, payload: ShareRequest,
                caller: Caller = Depends(get_caller),
                session: Session = Depends(get_session)) -> ShareOut:
    k = session.get(SchoolClass, class_id)
    if not k:
        raise HTTPException(status_code=404, detail="Класс не найден")
    if not _can_edit_class(session, k, caller.user_id):
        raise HTTPException(status_code=403, detail="Нет доступа к этому классу")
    existing = session.exec(
        select(ClassAccess).where(ClassAccess.class_id == class_id, ClassAccess.user_id == payload.user_id)
    ).first()
    if existing:
        existing.role = AccessRole(payload.role)
        session.add(existing)
    else:
        session.add(ClassAccess(class_id=class_id, user_id=payload.user_id, role=AccessRole(payload.role)))
    session.commit()
    return ShareOut(user_id=payload.user_id, role=AccessRole(payload.role))


@app.post("/api/v1/excursions", response_model=ExcursionOut, status_code=201, tags=["excursions"])
async def create_excursion(payload: ExcursionCreate, caller: Caller = Depends(get_caller),
                           session: Session = Depends(get_session)) -> ExcursionOut:
    klass = session.get(SchoolClass, payload.class_id)
    if not klass:
        raise HTTPException(status_code=404, detail="Класс не найден")
    event = session.get(CultureEvent, payload.culture_event_id) if payload.culture_event_id else None

    exc = Excursion(
        class_id=klass.id,
        culture_event_id=event.id if event else None,
        title=payload.title or (event.title if event else "Культурный выезд"),
        location_name=payload.location_name or (event.venue if event else ""),
        address=payload.address or (event.address if event else ""),
        event_date=payload.event_date or (event.event_date if event else None),
        gathering_time=_time_from_str(payload.gathering_time),
        return_time=_time_from_str(payload.return_time),
        deadline=_naive_utc(payload.deadline),
        ticket_price=payload.ticket_price if payload.ticket_price is not None else (event.price if event else 0.0),
        ticket_sale_url=payload.ticket_sale_url or (event.ticket_url if event else ""),
        is_pushkin_card=payload.is_pushkin_card if payload.is_pushkin_card is not None else (event.pushkin_eligible if event else False),
        status=ExcursionStatus.VOTING,
        school_name=klass.school_name,
        responsible_teacher=klass.teacher_name,
        order_basis=payload.order_basis or "план воспитательной работы",
        meeting_point=payload.meeting_point or "у главного входа школы",
        escorts=(payload.escorts or "").strip(),
        transport=(payload.transport or "").strip() or "пешком и на общественном транспорте",
        created_by_user_id=caller.user_id,
    )
    session.add(exc)
    session.commit()
    session.refresh(exc)

    for s in session.exec(select(Student).where(Student.class_id == klass.id)).all():
        session.add(ExcursionParticipant(
            excursion_id=exc.id, student_id=s.id,
            consent_status="PENDING",
            ticket_status=TicketStatus.NOT_REQUIRED if (exc.ticket_price or 0) == 0 else TicketStatus.WAITING_PAYMENT,
        ))
    session.commit()
    demo.schedule_virtual_answers(session, exc)
    if payload.notify_parents and settings.bot_enabled:
        try:
            await BotService().invite_parents(exc.id)
        except Exception:  # noqa: BLE001
            log.warning("Не удалось разослать приглашения по выезду %s", exc.id)
    return _excursion_out(exc, session, with_stats=True)


@app.get("/api/v1/excursions", response_model=list[ExcursionOut], tags=["excursions"])
def list_excursions(mine: bool = Query(default=False, description="Только выезды классов текущего пользователя"),
                    class_id: int | None = Query(default=None),
                    caller: Caller = Depends(get_caller),
                    session: Session = Depends(get_session)) -> list[ExcursionOut]:
    stmt = select(Excursion)
    if mine:
        stmt = stmt.where(Excursion.class_id.in_(_user_class_ids(session, caller.user_id) or [-1]))
    if class_id is not None:
        stmt = stmt.where(Excursion.class_id == class_id)
    rows = session.exec(stmt).all()
    rows = sorted(rows, key=lambda e: (e.event_date is None, e.event_date or date.max, e.id))
    return [_excursion_out(e, session, with_stats=mine or class_id is not None) for e in rows]


@app.get("/api/v1/excursions/{excursion_id}", response_model=ExcursionOut, tags=["excursions"])
def get_excursion(excursion_id: int, session: Session = Depends(get_session)) -> ExcursionOut:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    return _excursion_out(exc, session, with_stats=True)


@app.get("/api/v1/excursions/{excursion_id}/dashboard", response_model=DashboardOut, tags=["excursions"])
def excursion_dashboard(excursion_id: int, session: Session = Depends(get_session)) -> DashboardOut:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    data = services.dashboard(session, exc)
    return DashboardOut(
        excursion=_excursion_out(exc, session, with_stats=True),
        summary=data["summary"],
        progress_percent=data["progress_percent"],
        participants=[ParticipantRow(**row) for row in data["participants"]],
    )


@app.get("/api/v1/excursions/{excursion_id}/participants", tags=["excursions"])
def excursion_participants(excursion_id: int, session: Session = Depends(get_session)) -> dict:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    data = services.dashboard(session, exc)
    return {"excursion": _excursion_out(exc).model_dump(mode="json"), "participants": data["participants"]}


@app.post("/api/v1/excursions/{excursion_id}/consent", response_model=ConsentResponse, tags=["consent"])
async def post_consent(excursion_id: int, payload: ConsentRequest, request: Request,
                       session: Session = Depends(get_session)) -> ConsentResponse:
    _require_valid_parent(payload.init_data or request.headers.get("X-Max-Init-Data"))
    query_uid = request.query_params.get("user_id", "")
    caller = caller_from_request(request, payload.max_user_id or (int(query_uid) if query_uid.isdigit() else None))
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    if payload.status not in ("APPROVED", "REJECTED"):
        raise HTTPException(status_code=422, detail="status должен быть APPROVED или REJECTED")
    paper = payload.source == "paper"
    if caller.user_id:
        if paper:
            klass = session.get(SchoolClass, exc.class_id)
            if not (klass and _can_edit_class(session, klass, caller.user_id)):
                raise HTTPException(status_code=403, detail="Бумажные бланки отмечает только классный руководитель")
        elif not any(p.max_user_id == caller.user_id for p in services.student_parents(session, payload.student_id)):
            raise HTTPException(status_code=403, detail="Ответить может только родитель этого ученика")
        else:
            profile = session.exec(select(UserProfile).where(UserProfile.max_user_id == caller.user_id)).first()
            if not (profile and profile.terms_accepted_at):
                raise HTTPException(status_code=409, detail="Сначала примите соглашение о подписи "
                                                            "и дайте согласие на обработку данных")
    elif settings.MAX_VALIDATE_INIT_DATA or settings.bot_enabled:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX, чтобы ответить")
    try:
        p, changed, message = services.apply_consent(
            session, exc, student_id=payload.student_id, status=payload.status,
            parent_phone=payload.parent_phone, parent_name=payload.parent_name or caller.display_name,
            reason=payload.reason, source=payload.source,
            # Бумажный бланк — подпись родителя на бумаге; MAX ID учителя в подпись не записываем.
            max_user_id=None if paper else caller.user_id,
        )
    except ValueError as exc_err:
        raise HTTPException(status_code=404, detail=str(exc_err)) from exc_err
    if changed:
        try:
            bot = BotService()
            await bot.notify_consent(excursion_id, payload.student_id, exclude_user_id=caller.user_id)
            await bot.notify_teacher_if_complete(excursion_id)
        except Exception:  # noqa: BLE001
            log.warning("Не удалось уведомить второго родителя по выезду %s", excursion_id)
    return ConsentResponse(
        student_id=p.student_id, consent_status=p.consent_status, ticket_status=p.ticket_status,
        changed=changed, message=message, signed_by_name=p.signed_by_name, signed_at=p.signed_at,
    )


@app.post("/api/v1/excursions/{excursion_id}/consent/reset", response_model=ConsentResponse, tags=["consent"])
def reset_consent(excursion_id: int, payload: ConsentResetRequest, caller: Caller = Depends(get_caller),
                  session: Session = Depends(get_session)) -> ConsentResponse:
    """«Изменить ответ»: родитель — свой ответ до дедлайна, учитель — любую отметку своего класса."""
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    as_teacher = False
    if payload.as_teacher:
        klass = session.get(SchoolClass, exc.class_id)
        if not (klass and caller.user_id and _can_edit_class(session, klass, caller.user_id)):
            raise HTTPException(status_code=403, detail="Снять отметку может только классный руководитель")
        as_teacher = True
    try:
        p, changed, message = services.reset_consent(
            session, exc, student_id=payload.student_id, user_id=caller.user_id,
            as_teacher=as_teacher, source="miniapp",
        )
    except ValueError as exc_err:
        raise HTTPException(status_code=404, detail=str(exc_err)) from exc_err
    return ConsentResponse(
        student_id=p.student_id, consent_status=p.consent_status, ticket_status=p.ticket_status,
        changed=changed, message=message, signed_by_name=p.signed_by_name, signed_at=p.signed_at,
    )


def _teacher_excursion(session: Session, excursion_id: int, caller: Caller) -> Excursion:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    klass = session.get(SchoolClass, exc.class_id)
    if not klass or not _can_edit_class(session, klass, caller.user_id):
        raise HTTPException(status_code=403, detail="Это может сделать только классный руководитель")
    return exc


@app.post("/api/v1/excursions/{excursion_id}/attendance", tags=["trip-day"])
def mark_attendance(excursion_id: int, payload: AttendanceRequest, caller: Caller = Depends(get_caller),
                    session: Session = Depends(get_session)) -> dict:
    """Перекличка в день выезда: на месте / не пришёл (или снять отметку)."""
    exc = _teacher_excursion(session, excursion_id, caller)
    p = services.get_participant(session, exc.id, payload.student_id)
    if not p:
        raise HTTPException(status_code=404, detail="Ученик не найден в этом выезде")
    p.present = payload.present
    session.add(p)
    session.commit()
    return {"ok": True, "student_id": p.student_id, "present": p.present}


@app.post("/api/v1/excursions/{excursion_id}/milestone", response_model=MilestoneOut, tags=["trip-day"])
async def trip_milestone(excursion_id: int, payload: MilestoneRequest, caller: Caller = Depends(get_caller),
                         session: Session = Depends(get_session)) -> MilestoneOut:
    """«Группа выехала» / «Группа вернулась»: родители получают сообщение в MAX."""
    exc = _teacher_excursion(session, excursion_id, caller)
    if exc.cancelled:
        raise HTTPException(status_code=409, detail="Выезд отменён")
    if payload.kind == "returned" and not exc.departed_at:
        raise HTTPException(status_code=409, detail="Сначала отметьте, что группа выехала")
    absent = [session.get(Student, p.student_id).full_name for p in exc.participants
              if p.consent_status == "APPROVED" and p.present is False]
    already = exc.departed_at if payload.kind == "departed" else exc.returned_at
    if already:
        # Повторное нажатие не рассылает сообщения второй раз.
        return MilestoneOut(kind=payload.kind, at=already, notified=0, absent=absent)
    now = datetime.utcnow()
    if payload.kind == "departed":
        exc.departed_at = now
    else:
        exc.returned_at = now
    session.add(exc)
    session.commit()
    notified = await BotService().notify_milestone(exc.id, payload.kind) if settings.bot_enabled else 0
    return MilestoneOut(kind=payload.kind, at=now, notified=notified, absent=absent)


@app.post("/api/v1/excursions/{excursion_id}/cancel", tags=["excursions"])
async def cancel_trip(excursion_id: int, caller: Caller = Depends(get_caller),
                      session: Session = Depends(get_session)) -> dict:
    """Отмена выезда: родители, которые уже ответили или ждут, получают сообщение."""
    exc = _teacher_excursion(session, excursion_id, caller)
    if exc.cancelled:
        return {"ok": True, "notified": 0}
    exc.cancelled = True
    exc.status = ExcursionStatus.CLOSED
    session.add(exc)
    session.commit()
    notified = await BotService().notify_cancelled(exc.id) if settings.bot_enabled else 0
    return {"ok": True, "notified": notified}


@app.post("/api/v1/excursions/{excursion_id}/ticket-confirm", response_model=ConsentResponse, tags=["consent"])
async def post_ticket_confirm(excursion_id: int, payload: TicketConfirmRequest,
                              session: Session = Depends(get_session)) -> ConsentResponse:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    try:
        p, changed, message = services.confirm_ticket(
            session, exc, student_id=payload.student_id, ticket_number=payload.ticket_number,
            source=payload.source,
        )
    except ValueError as exc_err:
        raise HTTPException(status_code=404, detail=str(exc_err)) from exc_err
    if changed:
        try:
            await BotService().notify_teacher_if_complete(excursion_id)
        except Exception:  # noqa: BLE001
            log.warning("Не удалось уведомить учителя по выезду %s", excursion_id)
    return ConsentResponse(
        student_id=p.student_id, consent_status=p.consent_status, ticket_status=p.ticket_status,
        changed=changed, message=message, signed_by_name=p.signed_by_name, signed_at=p.signed_at,
    )


@app.post("/api/v1/excursions/{excursion_id}/remind-unconfirmed", response_model=RemindResponse, tags=["consent"])
async def remind_unconfirmed(excursion_id: int, session: Session = Depends(get_session)) -> RemindResponse:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    bot = BotService()
    targeted, delivered = await bot.remind(exc, session)
    return RemindResponse(targeted=targeted, delivered=len(delivered), recipients=delivered)


@app.post("/api/v1/excursions/{excursion_id}/publish", tags=["consent"])
async def publish_excursion(excursion_id: int, chat_id: int | None = Query(default=None),
                            session: Session = Depends(get_session)) -> dict:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    target_chat = chat_id if chat_id is not None else class_chat_id(session, exc)
    if target_chat is None:
        raise HTTPException(status_code=400, detail="Чат класса не указан.")
    bot = BotService()
    result = await bot.publish_to_chat(exc, target_chat)
    return {"ok": bool(result) or not settings.bot_enabled, "chat_id": target_chat,
            "delivered": bool(result)}


@app.post("/api/v1/excursions/{excursion_id}/tag-unactivated", tags=["consent"])
async def tag_unactivated(excursion_id: int, chat_id: int | None = Query(default=None),
                          session: Session = Depends(get_session)) -> dict:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    names = services.unactivated_children(session, exc)
    if not names:
        return {"ok": True, "tagged": 0, "names": [], "delivered": False}
    target_chat = chat_id if chat_id is not None else class_chat_id(session, exc)
    if target_chat is None:
        raise HTTPException(status_code=400, detail="Чат класса не указан.")
    bot = BotService()
    result = await bot.tag_unactivated(exc, target_chat, names)
    return {"ok": bool(result) or not settings.bot_enabled, "tagged": len(names),
            "names": names, "delivered": bool(result)}


@app.post("/api/v1/excursions/{excursion_id}/export-order", response_model=DocumentOut, tags=["documents"])
def generate_order(excursion_id: int, session: Session = Depends(get_session)) -> DocumentOut:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    last = session.exec(
        select(DocumentArtifact).where(DocumentArtifact.excursion_id == excursion_id)
    ).all()
    version = len(last) + 1
    filename = f"Приказ_выезд_{excursion_id}.docx"
    session.add(DocumentArtifact(excursion_id=excursion_id, kind="order", version=version, filename=filename))
    session.commit()
    base = settings.PUBLIC_BASE_URL.rstrip("/")
    return DocumentOut(
        excursion_id=excursion_id,
        docx_url=f"{base}/api/v1/excursions/{excursion_id}/export-order?fmt=docx",
        pdf_url=f"{base}/api/v1/excursions/{excursion_id}/export-order?fmt=pdf",
        version=version,
        generated_at=datetime.utcnow(),
    )


@app.get("/api/v1/excursions/{excursion_id}/export-order", tags=["documents"])
def download_order(excursion_id: int, fmt: str = Query(default="docx", pattern="^(docx|pdf)$"),
                   session: Session = Depends(get_session)) -> Response:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    if fmt == "pdf":
        content = build_order_pdf(session, exc)
        return Response(content=content, media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="order_{excursion_id}.pdf"'})
    content = build_order_docx(session, exc)
    return Response(
        content=content,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="order_{excursion_id}.docx"'},
    )


@app.post("/api/v1/students/{student_id}/link-code", response_model=LinkCodeOut, tags=["consent"])
def create_link_code(student_id: int, session: Session = Depends(get_session)) -> LinkCodeOut:
    student = session.get(Student, student_id)
    if not student:
        raise HTTPException(status_code=404, detail="Ученик не найден")
    code = secrets.token_urlsafe(8)
    parent = session.exec(select(ParentContact).where(ParentContact.student_id == student_id)).first()
    session.add(BotLinkCode(code=code, student_id=student_id, parent_phone=parent.phone_number if parent else ""))
    session.commit()
    deep_link = f"https://max.ru/{settings.MAX_BOT_USERNAME}?start=bind_{code}"
    return LinkCodeOut(code=code, deep_link=deep_link, student_id=student_id)


@app.get("/api/v1/parent/context", tags=["parent"])
def parent_context(max_user_id: int = Query(..., description="MAX user_id родителя (после привязки)"),
                   session: Session = Depends(get_session)) -> dict:
    children = services.parent_context(session, max_user_id)
    return {"max_user_id": max_user_id, "children": children, "found": bool(children)}


@app.get("/api/v1/parents/search", response_model=list[ParentSearchRow], tags=["parent"])
def parents_search(query: str = Query(..., min_length=2),
                   session: Session = Depends(get_session)) -> list[ParentSearchRow]:
    q = query.strip().lower()
    rows: list[ParentSearchRow] = []
    for parent in session.exec(select(ParentContact)).all():
        student = session.get(Student, parent.student_id)
        if not student:
            continue
        haystack = f"{parent.full_name} {parent.phone_number} {student.full_name}".lower()
        if q not in haystack:
            continue
        klass = session.get(SchoolClass, student.class_id)
        rows.append(ParentSearchRow(
            parent_id=parent.id, parent_name=parent.full_name, role=parent.role,
            phone=parent.phone_number, student_id=student.id, student_name=student.full_name,
            class_title=klass.title if klass else "", claimed=parent.max_user_id is not None,
        ))
        if len(rows) >= 20:
            break
    return rows


@app.post("/api/v1/parent/claim", tags=["parent"])
def parent_claim(payload: ParentClaimRequest, session: Session = Depends(get_session)) -> dict:
    parent = session.get(ParentContact, payload.parent_id)
    if not parent:
        raise HTTPException(status_code=404, detail="Контакт родителя не найден")
    for other in session.exec(
        select(ParentContact).where(ParentContact.max_user_id == payload.max_user_id)
    ).all():
        if other.id != parent.id and other.student_id == parent.student_id:
            other.max_user_id = None
            session.add(other)
    parent.max_user_id = payload.max_user_id
    parent.confirmed = True
    session.add(parent)
    session.commit()
    student = session.get(Student, parent.student_id)
    return {"ok": True, "parent_id": parent.id,
            "student_name": student.full_name if student else ""}


@app.post("/api/v1/parent/confirm-link", tags=["parent"])
def parent_confirm_link(payload: ParentLinkAction, session: Session = Depends(get_session)) -> dict:
    parent = services.confirm_parent_link(session, payload.student_id, payload.max_user_id, payload.accept)
    if not parent:
        raise HTTPException(status_code=404, detail="Запись не найдена")
    return {"ok": True, "confirmed": parent.confirmed,
            "linked": parent.max_user_id is not None}


@app.post("/api/v1/parent/revoke", tags=["parent"])
def parent_revoke(caller: Caller = Depends(get_caller), session: Session = Depends(get_session)) -> dict:
    """Отозвать согласие на обработку данных и отвечать на бумаге: отвязываем аккаунт MAX от детей."""
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    contacts = session.exec(select(ParentContact).where(ParentContact.max_user_id == caller.user_id)).all()
    for contact in contacts:
        contact.max_user_id = None
        contact.confirmed = False
        session.add(contact)
    profile = session.exec(select(UserProfile).where(UserProfile.max_user_id == caller.user_id)).first()
    if profile:
        profile.terms_accepted_at = None
        profile.role = None
        session.add(profile)
    session.commit()
    return {"ok": True, "unlinked": len(contacts)}


@app.post("/api/v1/parent/notifications", tags=["parent"])
def parent_notifications(payload: NotificationsRequest, session: Session = Depends(get_session)) -> dict:
    affected = services.set_notifications(session, payload.max_user_id, payload.enabled)
    return {"ok": True, "enabled": payload.enabled, "affected": affected}


def _children_count(session: Session, user_id: int | None) -> int:
    if not user_id:
        return 0
    return len(session.exec(select(ParentContact).where(ParentContact.max_user_id == user_id)).all())


@app.get("/api/v1/me", response_model=MeOut, tags=["me"])
def me(caller: Caller = Depends(get_caller), session: Session = Depends(get_session)) -> MeOut:
    """Кто открыл мини-приложение: имя из MAX, роли и есть ли демо-класс."""
    profile = remember_profile(session, caller)
    classes = _user_class_ids(session, caller.user_id)
    children = _children_count(session, caller.user_id)
    return MeOut(
        user_id=caller.user_id,
        name=(profile.display_name if profile else "") or caller.display_name,
        auth=caller.auth,
        role=profile.role if profile else None,
        is_teacher=bool(classes),
        is_parent=bool(children),
        has_demo=bool(caller.user_id and demo.sandbox_class(session, caller.user_id)),
        classes_count=len(classes),
        children_count=children,
        terms_accepted=bool(profile and profile.terms_accepted_at),
        bot_username=settings.MAX_BOT_USERNAME,
    )


@app.post("/api/v1/me/accept-terms", response_model=MeOut, tags=["me"])
def accept_terms(caller: Caller = Depends(get_caller), session: Session = Depends(get_session)) -> MeOut:
    """Родитель один раз принимает соглашение о ПЭП и даёт отдельное согласие на обработку ПДн."""
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    profile = remember_profile(session, caller)
    if not profile.terms_accepted_at:
        profile.terms_accepted_at = datetime.utcnow()
        session.add(profile)
        session.commit()
    return me(caller, session)


@app.post("/api/v1/me/role", response_model=MeOut, tags=["me"])
def set_role(payload: RoleRequest, caller: Caller = Depends(get_caller),
             session: Session = Depends(get_session)) -> MeOut:
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    profile = remember_profile(session, caller)
    profile.role = payload.role
    session.add(profile)
    session.commit()
    return me(caller, session)


@app.post("/api/v1/demo/start", response_model=DemoOut, tags=["demo"])
async def demo_start(caller: Caller = Depends(get_caller), session: Session = Depends(get_session)) -> DemoOut:
    """Личный демо-класс: пользователь — учитель 7Б и родитель одного ученика."""
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    profile = remember_profile(session, caller)
    info = demo.create_sandbox(session, caller.user_id, profile.display_name if profile else "")
    invited = 0
    if info["created"] and settings.bot_enabled:
        for excursion_id in info["excursion_ids"]:
            try:
                invited += await BotService().invite_parents(excursion_id)
            except Exception:  # noqa: BLE001
                log.warning("Не удалось отправить демо-приглашение по выезду %s", excursion_id)
    return DemoOut(**info, invited=invited)


@app.post("/api/v1/demo/reset", response_model=DemoOut, tags=["demo"])
async def demo_reset(caller: Caller = Depends(get_caller), session: Session = Depends(get_session)) -> DemoOut:
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    demo.delete_sandbox(session, caller.user_id)
    return await demo_start(caller, session)


def _invite_out(klass: SchoolClass) -> InviteOut:
    bot = settings.MAX_BOT_USERNAME
    return InviteOut(
        class_id=klass.id, code=klass.invite_code or "",
        app_link=f"https://max.ru/{bot}?startapp=join_{klass.invite_code}",
        bot_link=f"https://max.ru/{bot}?start=join_{klass.invite_code}",
    )


@app.post("/api/v1/classes/{class_id}/invite", response_model=InviteOut, tags=["classes"])
def class_invite(class_id: int, caller: Caller = Depends(get_caller),
                 session: Session = Depends(get_session)) -> InviteOut:
    """Одна ссылка (и QR) на весь класс: родитель открывает её и подтверждает номер."""
    klass = session.get(SchoolClass, class_id)
    if not klass:
        raise HTTPException(status_code=404, detail="Класс не найден")
    if not _can_edit_class(session, klass, caller.user_id):
        raise HTTPException(status_code=403, detail="Нет доступа к этому классу")
    if not klass.invite_code:
        klass.invite_code = demo.new_invite_code()
        session.add(klass)
        session.commit()
        session.refresh(klass)
    return _invite_out(klass)


async def _join_demo_class(session: Session, klass: SchoolClass, caller: Caller) -> JoinOut:
    profile = remember_profile(session, caller)
    name = (profile.display_name if profile else "") or caller.display_name
    student = demo.adopt_virtual_child(session, klass, caller.user_id, name)
    if not student:
        return JoinOut(status="not_found", message="В демо-классе не осталось свободных учеников.",
                       class_title=klass.title)
    if profile:
        profile.role = "parent"
        session.add(profile)
        session.commit()
    if settings.bot_enabled:
        bot = BotService()
        try:
            for exc in session.exec(select(Excursion).where(Excursion.class_id == klass.id)).all():
                await bot.invite_one(exc.id, student.id, caller.user_id)
            await bot.notify_owner_joined(klass, name or "Новый родитель", student.full_name)
        except Exception:  # noqa: BLE001
            log.warning("Не удалось отправить приглашение новому родителю демо-класса %s", klass.id)
    return JoinOut(status="linked", message="Готово: вы родитель ученика демо-класса.", children=[student.full_name],
                   class_title=klass.title)


@app.get("/api/v1/join/{code}", response_model=JoinInfoOut, tags=["parent"])
def join_info(code: str, session: Session = Depends(get_session)) -> JoinInfoOut:
    """Что за класс по ссылке-приглашению — показываем родителю до подтверждения номера."""
    klass = session.exec(select(SchoolClass).where(SchoolClass.invite_code == code)).first()
    if not klass:
        raise HTTPException(status_code=404, detail="Ссылка-приглашение не найдена")
    count = len(session.exec(select(Student).where(Student.class_id == klass.id)).all())
    return JoinInfoOut(code=code, is_demo=klass.is_demo, class_title=klass.title, school_name=klass.school_name,
                       teacher_name=klass.teacher_name, students_count=count)


def _pair_link(code: str) -> tuple[str, str]:
    bot = settings.MAX_BOT_USERNAME
    return f"https://max.ru/{bot}?startapp=pair_{code}", f"https://max.ru/{bot}?start=pair_{code}"


@app.post("/api/v1/parent/second-invite", response_model=SecondParentInviteOut, tags=["parent"])
def second_parent_invite(payload: SecondParentInviteRequest, caller: Caller = Depends(get_caller),
                         session: Session = Depends(get_session)) -> SecondParentInviteOut:
    """«Позвать второго родителя»: ссылка, по которой ещё один взрослый подключится к ребёнку.

    Пересылка сообщения бота не подходит — кнопки при пересылке пропадают, поэтому даём ссылку.
    """
    student = session.get(Student, payload.student_id)
    if not student:
        raise HTTPException(status_code=404, detail="Ученик не найден")
    bound = session.exec(select(ParentContact).where(
        ParentContact.student_id == student.id, ParentContact.max_user_id == caller.user_id)).first()
    if not caller.user_id or not bound:
        raise HTTPException(status_code=403, detail="Позвать второго родителя может только подключённый родитель")
    code = secrets.token_urlsafe(9).replace("-", "").replace("_", "")[:12]
    session.add(BotLinkCode(code=code, student_id=student.id, new_parent=True, invited_by_user_id=caller.user_id))
    session.commit()
    app_link, bot_link = _pair_link(code)
    return SecondParentInviteOut(code=code, app_link=app_link, bot_link=bot_link, student_name=student.full_name)


@app.get("/api/v1/parent/pair/{code}", response_model=PairInfoOut, tags=["parent"])
def pair_info(code: str, session: Session = Depends(get_session)) -> PairInfoOut:
    link = session.exec(select(BotLinkCode).where(BotLinkCode.code == code, BotLinkCode.new_parent == True)).first()  # noqa: E712
    if not link or link.used:
        raise HTTPException(status_code=404, detail="Ссылка уже использована или устарела")
    student = session.get(Student, link.student_id)
    klass = session.get(SchoolClass, student.class_id) if student else None
    inviter = session.exec(select(ParentContact).where(
        ParentContact.student_id == link.student_id, ParentContact.max_user_id == link.invited_by_user_id)).first()
    return PairInfoOut(code=code, student_name=student.full_name if student else "", class_title=klass.title if klass else "",
                       invited_by=inviter.full_name if inviter else "")


@app.post("/api/v1/parent/pair", response_model=JoinOut, tags=["parent"])
def accept_pair(payload: PairRequest, caller: Caller = Depends(get_caller),
                session: Session = Depends(get_session)) -> JoinOut:
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    profile = remember_profile(session, caller)
    student = services.accept_pair_code(session, payload.code, caller.user_id,
                                        profile.display_name if profile else caller.display_name)
    if not student:
        return JoinOut(status="bad_code", message="Ссылка уже использована или устарела. Попросите прислать новую.")
    if profile:
        profile.role = "parent"
        session.add(profile)
        session.commit()
    klass = session.get(SchoolClass, student.class_id)
    return JoinOut(status="linked", message="Готово: вы подключены как второй родитель.",
                   children=[student.full_name], class_title=klass.title if klass else None)


@app.post("/api/v1/parent/join", response_model=JoinOut, tags=["parent"])
async def parent_join(payload: JoinRequest, caller: Caller = Depends(get_caller),
                      session: Session = Depends(get_session)) -> JoinOut:
    """Привязка родителя по номеру телефона, подтверждённому в MAX.

    Номер сверяется со списком класса, который загрузил учитель. Совпал —
    родитель сразу видит приглашения своего ребёнка. Без выбора «себя» из базы.
    Ссылка демо-класса работает без телефона: открывший становится родителем
    свободного ученика (живое демо на защите).
    """
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    if payload.code:
        demo_class = session.exec(select(SchoolClass).where(
            SchoolClass.invite_code == payload.code, SchoolClass.is_demo == True)).first()  # noqa: E712
        if demo_class:
            return await _join_demo_class(session, demo_class, caller)
    if caller.verified and settings.MAX_BOT_TOKEN:
        # Внутри MAX номер принимаем только с подписью платформы.
        if not verify_phone(payload.auth_date or "", payload.phone, caller.user_id, payload.hash or "",
                            settings.MAX_BOT_TOKEN):
            raise HTTPException(status_code=403, detail="MAX не подтвердил номер. Попробуйте ещё раз.")
    elif settings.MAX_VALIDATE_INIT_DATA or settings.bot_enabled:
        # На боевом стенде номер принимаем только подписанным MAX — ввести чужой нельзя.
        raise HTTPException(status_code=401, detail="Откройте ссылку в приложении MAX.")
    klass = None
    if payload.code:
        klass = session.exec(select(SchoolClass).where(SchoolClass.invite_code == payload.code)).first()
        if not klass:
            return JoinOut(status="bad_code", message="Ссылка-приглашение не найдена. Попросите у учителя новую.")
    phone = services.normalize_phone(payload.phone)
    if not phone:
        return JoinOut(status="no_phone", message="Подтвердите номер телефона, чтобы найти вашего ребёнка.",
                       class_title=klass.title if klass else None)
    stmt = select(ParentContact)
    if klass:
        student_ids = [s.id for s in session.exec(select(Student).where(Student.class_id == klass.id)).all()]
        stmt = stmt.where(ParentContact.student_id.in_(student_ids or [-1]))
    matches = [p for p in session.exec(stmt).all()
               if not p.is_virtual and services.normalize_phone(p.phone_number) == phone]
    if not matches:
        return JoinOut(
            status="not_found",
            message="Этого номера нет в списке класса. Попросите классного руководителя "
                    "добавить его и откройте ссылку ещё раз.",
            class_title=klass.title if klass else None,
        )
    names = []
    profile = remember_profile(session, caller)
    for parent in matches:
        parent.max_user_id = caller.user_id
        parent.confirmed = True
        if profile and profile.display_name and not parent.full_name:
            parent.full_name = profile.display_name
        session.add(parent)
        student = session.get(Student, parent.student_id)
        if student:
            names.append(student.full_name)
    if profile:
        profile.role = "parent"
        session.add(profile)
    session.commit()
    return JoinOut(status="linked", message="Готово: вы подключены к классу.", children=names,
                   class_title=klass.title if klass else None)


@app.post("/api/v1/excursions/{excursion_id}/send-documents", response_model=SendDocumentsOut, tags=["documents"])
async def send_documents(excursion_id: int, caller: Caller = Depends(get_caller),
                         session: Session = Depends(get_session)) -> SendDocumentsOut:
    """Приказ, список, реестр согласий и инструктаж — файлами в чат MAX учителя."""
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    if not caller.user_id:
        raise HTTPException(status_code=401, detail="Откройте ClassGo в MAX")
    if not settings.bot_enabled:
        return SendDocumentsOut(delivered=False, message="Отправка в чат сейчас не работает. Скачайте PDF или DOCX.")
    delivered = await BotService().send_documents(session, exc, caller.user_id)
    return SendDocumentsOut(
        delivered=delivered,
        message="Документы отправлены вам в чат с ботом." if delivered
        else "Не получилось отправить файлы в чат. Скачайте PDF или DOCX.",
    )


@app.post("/api/v1/excursions/{excursion_id}/invite-parents", tags=["consent"])
async def invite_parents(excursion_id: int, session: Session = Depends(get_session)) -> dict:
    exc = session.get(Excursion, excursion_id)
    if not exc:
        raise HTTPException(status_code=404, detail="Выезд не найден")
    sent = await BotService().invite_parents(excursion_id) if settings.bot_enabled else 0
    return {"ok": True, "sent": sent}


@app.post("/api/v1/max/validate-init-data", tags=["max"])
def validate_init_data(payload: dict) -> dict:
    init_data = (payload or {}).get("init_data", "")
    if not settings.MAX_BOT_TOKEN:
        return {"valid": False, "reason": "bot_token_not_configured"}
    valid, fields = validate_webapp_data(init_data, settings.MAX_BOT_TOKEN)
    return {
        "valid": valid,
        "fresh": is_fresh(fields.get("auth_date")),
        "user": fields.get("user"),
        "chat": fields.get("chat"),
    }


@app.post("/api/v1/admin/reset-demo", tags=["service"])
def reset_demo(session: Session = Depends(get_session)) -> dict:
    exc = session.exec(select(Excursion)).first()
    if not exc:
        raise HTTPException(status_code=404, detail="Демо-выезд не создан")
    changed = 0
    for p in session.exec(select(ExcursionParticipant).where(ExcursionParticipant.excursion_id == exc.id)).all():
        p.consent_status = "PENDING"
        p.ticket_status = TicketStatus.WAITING_PAYMENT if (exc.ticket_price or 0) > 0 else TicketStatus.NOT_REQUIRED
        p.signed_at = None
        p.signed_by_name = None
        p.signed_by_phone = None
        p.rejection_reason = None
        p.ticket_number = None
        session.add(p)
        changed += 1
    session.commit()
    return {"excursion_id": exc.id, "reset_participants": changed}


@app.post("/webhook/max", tags=["max"])
async def max_webhook(request: Request) -> JSONResponse:
    try:
        payload = await request.json()
    except Exception:  # noqa: BLE001
        return JSONResponse({"ok": False, "error": "invalid json"}, status_code=400)
    updates = payload if isinstance(payload, list) else [payload]
    bot = BotService()
    processed = 0
    for update in updates:
        if isinstance(update, dict) and update.get("update_type"):
            await bot.handle_update(update)
            processed += 1
    return JSONResponse({"ok": True, "processed": processed})
