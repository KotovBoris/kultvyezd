"""Пакет документов по выезду — по образцам письма Минпросвещения России от 16.06.2022 № 06-836.

Состав (один PDF для печати и такой же DOCX для правки под шапку школы):
  • проект приказа об организации выездного мероприятия (по полям прил. 1 к 06-836);
  • Приложение 1 — список группы с контактами законных представителей;
  • Приложение 2 — реестр согласий: кто, когда и каким способом ответил (ПЭП, 63-ФЗ);
  • Приложение 3 — ведомость целевого инструктажа (по форме прил. 3) и памятка (прил. 2);
  • бумажные бланки согласия для тех, кто ещё не ответил или не пользуется MAX.
"""
from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import datetime, timedelta

from docx import Document
from docx.enum.section import WD_ORIENT  # noqa: F401  (оставлено для ручной правки шаблона)
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from sqlmodel import Session, select

from .models import ConsentStatus, Excursion, ExcursionParticipant, ParentContact, SchoolClass, Student

MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
          "сентября", "октября", "ноября", "декабря"]

SOURCE_LABEL = {
    "bot": "MAX, кнопка в чате",
    "miniapp": "MAX, приложение ClassGo",
    "demo": "MAX (демо)",
    "paper": "бумажное согласие",
    "call": "по телефону",
    "person": "лично",
    "api": "ClassGo",
}

MEMO = [
    "Слушать руководителя группы и сопровождающих, выполнять их указания.",
    "Не отходить от группы и не покидать место проведения мероприятия без разрешения руководителя.",
    "Соблюдать правила дорожного движения: переходить дорогу только по пешеходному переходу и вместе с группой.",
    "В общественном транспорте держаться за поручни, не стоять у дверей, не мешать другим пассажирам.",
    "В музее и театре соблюдать правила учреждения: не трогать экспонаты без разрешения, не шуметь.",
    "При плохом самочувствии или происшествии сразу сообщить руководителю группы.",
    "Если отстал от группы — оставаться на месте и позвонить руководителю группы.",
    "Иметь при себе заряженный телефон; номер руководителя группы записан в контактах.",
]


@dataclass
class Row:
    name: str
    birth: str
    parent: str
    phone: str
    status: ConsentStatus
    reason: str = ""
    signed_by: str = ""
    signed_at: str = ""
    method: str = ""
    max_id: str = ""
    ticket: str = ""


@dataclass
class Pack:
    school: str
    class_title: str
    title: str
    place: str
    date_long: str
    date_short: str
    gathering: str
    returning: str
    meeting_point: str
    teacher: str
    escorts: str
    basis: str
    deadline: str
    going: list[Row] = field(default_factory=list)
    not_going: list[Row] = field(default_factory=list)
    waiting: list[Row] = field(default_factory=list)


def _date_long(value) -> str:
    if not value:
        return "«___» ____________ 20__ г."
    return f"«{value.day}» {MONTHS[value.month - 1]} {value.year} г."


def _msk(value: datetime | None) -> str:
    """Время хранится в UTC — в документах показываем московское."""
    if not value:
        return "—"
    return (value + timedelta(hours=3)).strftime("%d.%m.%Y %H:%M")


def _quoted(title: str) -> str:
    return title if "«" in title else f"«{title}»"


def _basis(text: str) -> str:
    """«план воспитательной работы» → «планом воспитательной работы» для преамбулы приказа."""
    text = (text or "план воспитательной работы").strip()
    if text.lower().startswith("план "):
        return "планом " + text[5:]
    return text


def _collect(session: Session, excursion: Excursion) -> Pack:
    klass = session.get(SchoolClass, excursion.class_id)
    parts = session.exec(
        select(ExcursionParticipant).where(ExcursionParticipant.excursion_id == excursion.id)
    ).all()
    pack = Pack(
        school=excursion.school_name or (klass.school_name if klass else "") or "Образовательная организация",
        class_title=klass.title if klass else "",
        title=_quoted(excursion.title),
        place=", ".join(p for p in (excursion.location_name, excursion.address) if p),
        date_long=_date_long(excursion.event_date),
        date_short=excursion.event_date.strftime("%d.%m.%Y") if excursion.event_date else "__.__.____",
        gathering=excursion.gathering_time.strftime("%H:%M") if excursion.gathering_time else "__:__",
        returning=excursion.return_time.strftime("%H:%M") if excursion.return_time else "__:__",
        meeting_point=excursion.meeting_point or "у главного входа школы",
        teacher=excursion.responsible_teacher or (klass.teacher_name if klass else "") or "",
        escorts=excursion.escorts or "",
        basis=_basis(excursion.order_basis),
        deadline=_msk(excursion.deadline),
    )
    for p in parts:
        student = session.get(Student, p.student_id)
        parents = session.exec(select(ParentContact).where(ParentContact.student_id == p.student_id)).all()
        parent = next((x for x in parents if x.max_user_id), parents[0] if parents else None)
        row = Row(
            name=student.full_name if student else f"#{p.student_id}",
            birth=student.birth_date.strftime("%d.%m.%Y") if student and student.birth_date else "—",
            parent=f"{parent.full_name} ({parent.role.lower()})" if parent and parent.full_name else "—",
            phone=parent.phone_number if parent and parent.phone_number else "—",
            status=p.consent_status,
            reason=p.rejection_reason or "",
            signed_by=p.signed_by_name or "",
            signed_at=_msk(p.signed_at),
            method=SOURCE_LABEL.get(p.signed_source or "", "ClassGo"),
            max_id=str(p.signed_by_max_user_id) if p.signed_by_max_user_id else "—",
            ticket={"PAID": "куплен", "WAITING_PAYMENT": "не отмечен", "NOT_REQUIRED": "не нужен"}.get(
                p.ticket_status.value, ""),
        )
        if p.consent_status == ConsentStatus.APPROVED:
            pack.going.append(row)
        elif p.consent_status == ConsentStatus.REJECTED:
            pack.not_going.append(row)
        else:
            pack.waiting.append(row)
    for lst in (pack.going, pack.not_going, pack.waiting):
        lst.sort(key=lambda r: r.name)
    return pack


def _order_points(pack: Pack) -> list[str]:
    n = len(pack.going)
    escorts = pack.escorts or "________________________________"
    return [
        f"Организовать {pack.date_short} выездное экскурсионное мероприятие {pack.title} "
        f"для обучающихся {pack.class_title} класса в количестве {n} чел. согласно списку (приложение 1). "
        f"Место проведения: {pack.place}.",
        f"Назначить руководителем группы {pack.teacher or '________________________'}, "
        f"возложив на него(неё) ответственность за жизнь и здоровье обучающихся во время мероприятия. "
        f"Назначить сопровождающим (помощником руководителя группы): {escorts}.",
        f"Установить время сбора группы — {pack.gathering} {pack.meeting_point}; "
        f"контрольное время возвращения — {pack.returning}. О возвращении группы руководителю "
        f"сообщить ответственному за контроль.",
        "Руководителю группы: провести целевой инструктаж обучающихся по правилам безопасности под подпись "
        "(приложение 3); иметь при себе список группы с контактами законных представителей (приложение 1); "
        "обеспечить наличие согласий законных представителей на участие обучающихся (приложение 2).",
        "Контроль за исполнением приказа возложить на заместителя директора по воспитательной работе "
        "________________________.",
    ]


PREAMBLE = (
    "В соответствии с Федеральным законом от 29.12.2012 № 273-ФЗ «Об образовании в Российской Федерации», "
    "методическими рекомендациями Минпросвещения России по организации экскурсий и выездных мероприятий "
    "(письмо от 16.06.2022 № 06-836) и {basis}, в целях развития познавательного интереса обучающихся"
)


# ---------------------------------------------------------------------------- PDF
_FONTS: tuple[str, str] | None = None


def _fonts() -> tuple[str, str]:
    global _FONTS
    if _FONTS:
        return _FONTS
    pairs = [
        ("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf", "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"),
        ("/usr/share/fonts/dejavu/DejaVuSans.ttf", "/usr/share/fonts/dejavu/DejaVuSans-Bold.ttf"),
        ("C:/Windows/Fonts/arial.ttf", "C:/Windows/Fonts/arialbd.ttf"),
        ("/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Supplemental/Arial Bold.ttf"),
    ]
    for regular, bold in pairs:
        try:
            pdfmetrics.registerFont(TTFont("CG", regular))
            try:
                pdfmetrics.registerFont(TTFont("CG-Bold", bold))
            except Exception:  # noqa: BLE001
                pdfmetrics.registerFont(TTFont("CG-Bold", regular))
            # Семейство нужно, чтобы <b> в абзацах переключался на жирное начертание.
            pdfmetrics.registerFontFamily("CG", normal="CG", bold="CG-Bold", italic="CG", boldItalic="CG-Bold")
            _FONTS = ("CG", "CG-Bold")
            return _FONTS
        except Exception:  # noqa: BLE001
            continue
    _FONTS = ("Helvetica", "Helvetica-Bold")
    return _FONTS


def _styles() -> dict[str, ParagraphStyle]:
    font, bold = _fonts()
    base = ParagraphStyle("base", fontName=font, fontSize=11, leading=15)
    return {
        "school": ParagraphStyle("school", parent=base, fontName=bold, alignment=TA_CENTER, fontSize=11.5),
        "draft": ParagraphStyle("draft", parent=base, alignment=TA_CENTER, fontSize=9, textColor=colors.grey),
        "h1": ParagraphStyle("h1", parent=base, fontName=bold, fontSize=14, leading=18, alignment=TA_CENTER,
                             spaceBefore=6, spaceAfter=2),
        "h2": ParagraphStyle("h2", parent=base, fontName=bold, fontSize=12, leading=16, spaceBefore=4, spaceAfter=6),
        "center": ParagraphStyle("center", parent=base, alignment=TA_CENTER),
        "body": ParagraphStyle("body", parent=base, alignment=TA_JUSTIFY, firstLineIndent=10 * mm),
        "item": ParagraphStyle("item", parent=base, alignment=TA_LEFT, spaceBefore=3),
        "small": ParagraphStyle("small", parent=base, fontSize=9, leading=12, textColor=colors.HexColor("#444444")),
        "cell": ParagraphStyle("cell", parent=base, fontSize=9.5, leading=12),
        "cellb": ParagraphStyle("cellb", parent=base, fontName=bold, fontSize=9.5, leading=12),
        "appx": ParagraphStyle("appx", parent=base, fontSize=9.5, leading=12, alignment=2),
    }


def _table(data: list[list], widths: list[float], s: dict[str, ParagraphStyle]) -> Table:
    rows = [[Paragraph(str(c), s["cellb"] if i == 0 else s["cell"]) for c in row] for i, row in enumerate(data)]
    t = Table(rows, colWidths=widths, repeatRows=1)
    t.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#9aa1ab")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eef1f5")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return t


def _sign_line(label: str, s: dict[str, ParagraphStyle]) -> Paragraph:
    return Paragraph(f"{label} ______________________ / ______________________", s["item"])


def build_order_pdf(session: Session, excursion: Excursion) -> bytes:
    pack = _collect(session, excursion)
    s = _styles()
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=25 * mm, rightMargin=15 * mm,
                            topMargin=15 * mm, bottomMargin=15 * mm,
                            title=f"Приказ о выезде {pack.title}", author="ClassGo")
    story: list = []

    # --- Проект приказа
    story += [
        Paragraph(pack.school, s["school"]),
        Paragraph("проект документа подготовлен в ClassGo", s["draft"]),
        Spacer(1, 6 * mm),
        Paragraph("ПРИКАЗ", s["h1"]),
        Paragraph("№ ________ &nbsp;&nbsp;&nbsp; от «___» ____________ 20__ г.", s["center"]),
        Spacer(1, 4 * mm),
        Paragraph(f"<b>Об организации выездного экскурсионного мероприятия обучающихся {pack.class_title} класса</b>",
                  s["center"]),
        Spacer(1, 5 * mm),
        Paragraph(PREAMBLE.format(basis=pack.basis) + ",", s["body"]),
        Spacer(1, 2 * mm),
        Paragraph("<b>ПРИКАЗЫВАЮ:</b>", s["item"]),
    ]
    for i, text in enumerate(_order_points(pack), 1):
        story.append(Paragraph(f"{i}. {text}", s["item"]))
    story += [
        Spacer(1, 10 * mm),
        _sign_line("Директор", s),
        Spacer(1, 6 * mm),
        Paragraph("С приказом ознакомлены:", s["item"]),
        Paragraph(f"Руководитель группы: {pack.teacher or '________________'} &nbsp; ____________ (подпись)", s["item"]),
        _sign_line("Сопровождающий", s),
        PageBreak(),
    ]

    # --- Приложение 1. Список группы
    story += [
        Paragraph(f"Приложение 1 к приказу от «___» ________ 20__ г. № ____", s["appx"]),
        Paragraph(f"Список группы обучающихся {pack.class_title} класса", s["h2"]),
        Paragraph(f"Мероприятие: {pack.title}, {pack.date_short}. Место: {pack.place}.", s["small"]),
        Spacer(1, 3 * mm),
    ]
    data = [["№", "ФИО обучающегося", "Дата рождения", "Законный представитель", "Телефон"]]
    for i, r in enumerate(pack.going, 1):
        data.append([i, r.name, r.birth, r.parent, r.phone])
    if not pack.going:
        data.append(["", "Пока никто не отпущен", "", "", ""])
    story.append(_table(data, [9 * mm, 48 * mm, 22 * mm, 55 * mm, 33 * mm], s))
    story += [
        Spacer(1, 4 * mm),
        Paragraph(f"Всего обучающихся: {len(pack.going)}. Руководитель группы: {pack.teacher or '________'}. "
                  f"Сопровождающие: {pack.escorts or '________'}.", s["item"]),
        Paragraph("Дата рождения указывается при организованной перевозке группы детей автобусом "
                  "(ПП РФ от 02.06.2026 № 680, п. 13).", s["small"]),
        PageBreak(),
    ]

    # --- Приложение 2. Реестр согласий
    story += [
        Paragraph("Приложение 2 к приказу от «___» ________ 20__ г. № ____", s["appx"]),
        Paragraph("Реестр согласий законных представителей", s["h2"]),
        Paragraph(f"Ответы собраны в ClassGo до {pack.deadline} (МСК).", s["small"]),
        Spacer(1, 3 * mm),
    ]
    data = [["№", "Обучающийся", "Ответ и кто ответил", "Способ", "Дата и время (МСК)", "MAX ID"]]
    rows = [(r, "Отпускаю") for r in pack.going] + [(r, f"Не сможет: {r.reason or 'без причины'}") for r in pack.not_going]
    for i, (r, verdict) in enumerate(rows, 1):
        data.append([i, r.name, f"{verdict}<br/>{r.signed_by or '—'}", r.method, r.signed_at, r.max_id])
    if not rows:
        data.append(["", "Ответов пока нет", "", "", "", ""])
    story.append(_table(data, [8 * mm, 38 * mm, 48 * mm, 28 * mm, 25 * mm, 20 * mm], s))
    if pack.waiting:
        story += [
            Spacer(1, 3 * mm),
            Paragraph("Ещё не ответили: " + ", ".join(r.name for r in pack.waiting)
                      + ". Для них в пакет добавлены бумажные бланки согласия.", s["small"]),
        ]
    story += [
        Spacer(1, 3 * mm),
        Paragraph("Электронные ответы подписаны простой электронной подписью законного представителя "
                  "(ст. 160 ГК РФ; ст. 6, 9 Федерального закона от 06.04.2011 № 63-ФЗ): лицо определяется по "
                  "аккаунту MAX и подтверждённому платформой номеру телефона, соглашение о ПЭП принято в ClassGo. "
                  "Ответы, отмеченные учителем, получены вне MAX (бумага, телефон, лично).", s["small"]),
        PageBreak(),
    ]

    # --- Приложение 3. Ведомость инструктажа + памятка
    story += [
        Paragraph("Приложение 3 к приказу от «___» ________ 20__ г. № ____", s["appx"]),
        Paragraph("Ведомость проведения целевого инструктажа по правилам безопасности", s["h2"]),
        Paragraph(f"Мероприятие: {pack.title}, {pack.date_short}, {pack.place}.", s["small"]),
        Spacer(1, 3 * mm),
    ]
    data = [["№ п/п", "ФИО инструктируемого", "Класс", "Подпись"]]
    for i, r in enumerate(pack.going, 1):
        data.append([i, r.name, pack.class_title, ""])
    for _ in range(max(0, 3 - len(pack.going))):
        data.append(["", "", "", ""])
    story.append(_table(data, [14 * mm, 82 * mm, 20 * mm, 52 * mm], s))
    story += [
        Spacer(1, 4 * mm),
        Paragraph(f"Присутствовало: ______ чел. Инструктаж провёл(а): {pack.teacher or '________________'} "
                  "____________ (подпись) &nbsp; Дата: ___________", s["item"]),
        Spacer(1, 6 * mm),
        Paragraph("Памятка обучающимся (содержание инструктажа)", s["h2"]),
    ]
    for i, line in enumerate(MEMO, 1):
        story.append(Paragraph(f"{i}. {line}", s["item"]))

    # --- Бумажные бланки для тех, кто ещё не ответил
    if pack.waiting:
        story.append(PageBreak())
        story.append(Paragraph("Бланки согласия для законных представителей, которые ещё не ответили", s["h2"]))
        story.append(Paragraph("Распечатайте и передайте родителям, которые не пользуются MAX. После возврата бланка "
                               "отметьте ответ в ClassGo: «Отметить вручную → бумажное согласие».", s["small"]))
        for r in pack.waiting:
            blank = [
                Spacer(1, 6 * mm),
                Paragraph("<b>Согласие законного представителя на участие обучающегося в выездном мероприятии</b>",
                          s["item"]),
                Paragraph(f"Я, ______________________________________________ (ФИО), законный представитель "
                          f"обучающегося <b>{r.name}</b>, {pack.class_title} класс, даю согласие на его (её) участие "
                          f"в выездном мероприятии {pack.title} {pack.date_short} ({pack.place}). Сбор в "
                          f"{pack.gathering} {pack.meeting_point}, возвращение к {pack.returning}. "
                          f"Руководитель группы — {pack.teacher or '________________'}.", s["item"]),
                Paragraph("Контактный телефон: ____________________ &nbsp;&nbsp; "
                          "Подпись: ____________ &nbsp;&nbsp; Дата: ____________", s["item"]),
                Paragraph("— — — — — — — — — — — — — — — — — — — — — — — — — — — — — — —", s["draft"]),
            ]
            story.append(KeepTogether(blank))

    def footer(canvas, doc_):
        font, _ = _fonts()
        canvas.saveState()
        canvas.setFont(font, 8)
        canvas.setFillColor(colors.grey)
        canvas.drawString(25 * mm, 8 * mm, f"ClassGo · {pack.class_title} · {pack.title} · {pack.date_short}")
        canvas.drawRightString(A4[0] - 15 * mm, 8 * mm, f"стр. {doc_.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buf.getvalue()


# ---------------------------------------------------------------------------- DOCX
def _docx_table(doc, header: list[str], rows: list[list]) -> None:
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    for i, h in enumerate(header):
        cell = t.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h).bold = True
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)


def _page_break(doc) -> None:
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def build_order_docx(session: Session, excursion: Excursion) -> bytes:
    pack = _collect(session, excursion)
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)

    def centered(text: str, bold: bool = False, size: int | None = None):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = bold
        if size:
            run.font.size = Pt(size)
        return p

    centered(pack.school, bold=True)
    centered("ПРИКАЗ", bold=True, size=14)
    centered("№ ________ от «___» ____________ 20__ г.")
    centered(f"Об организации выездного экскурсионного мероприятия обучающихся {pack.class_title} класса", bold=True)
    doc.add_paragraph(PREAMBLE.format(basis=pack.basis) + ",")
    doc.add_paragraph().add_run("ПРИКАЗЫВАЮ:").bold = True
    for i, text in enumerate(_order_points(pack), 1):
        doc.add_paragraph(f"{i}. {text}")
    doc.add_paragraph("\nДиректор ______________________ / ______________________")
    doc.add_paragraph("С приказом ознакомлены:")
    doc.add_paragraph(f"Руководитель группы {pack.teacher} ______________________")
    doc.add_paragraph("Сопровождающий ______________________ / ______________________")

    _page_break(doc)
    doc.add_paragraph("Приложение 1 к приказу от «___» ________ 20__ г. № ____").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph().add_run(f"Список группы обучающихся {pack.class_title} класса").bold = True
    doc.add_paragraph(f"Мероприятие: {pack.title}, {pack.date_short}. Место: {pack.place}.")
    _docx_table(doc, ["№", "ФИО обучающегося", "Дата рождения", "Законный представитель", "Телефон"],
                [[i, r.name, r.birth, r.parent, r.phone] for i, r in enumerate(pack.going, 1)])
    doc.add_paragraph(f"Всего обучающихся: {len(pack.going)}. Руководитель группы: {pack.teacher}. "
                      f"Сопровождающие: {pack.escorts or '________'}.")

    _page_break(doc)
    doc.add_paragraph("Приложение 2 к приказу от «___» ________ 20__ г. № ____").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph().add_run("Реестр согласий законных представителей").bold = True
    rows = [(r, "Отпускаю") for r in pack.going] + [(r, f"Не сможет: {r.reason or 'без причины'}") for r in pack.not_going]
    _docx_table(doc, ["№", "Обучающийся", "Ответ", "Кто ответил", "Способ", "Дата и время (МСК)", "MAX ID"],
                [[i, r.name, verdict, r.signed_by or "—", r.method, r.signed_at, r.max_id]
                 for i, (r, verdict) in enumerate(rows, 1)])
    if pack.waiting:
        doc.add_paragraph("Ещё не ответили: " + ", ".join(r.name for r in pack.waiting) + ".")
    doc.add_paragraph("Электронные ответы подписаны простой электронной подписью законного представителя "
                      "(ст. 160 ГК РФ; ст. 6, 9 Федерального закона № 63-ФЗ).")

    _page_break(doc)
    doc.add_paragraph("Приложение 3 к приказу от «___» ________ 20__ г. № ____").alignment = WD_ALIGN_PARAGRAPH.RIGHT
    doc.add_paragraph().add_run("Ведомость проведения целевого инструктажа по правилам безопасности").bold = True
    _docx_table(doc, ["№ п/п", "ФИО инструктируемого", "Класс", "Подпись"],
                [[i, r.name, pack.class_title, ""] for i, r in enumerate(pack.going, 1)])
    doc.add_paragraph(f"Присутствовало: ______ чел. Инструктаж провёл(а): {pack.teacher} ____________ "
                      "Дата: ___________")
    doc.add_paragraph().add_run("Памятка обучающимся").bold = True
    for i, line in enumerate(MEMO, 1):
        doc.add_paragraph(f"{i}. {line}")

    if pack.waiting:
        _page_break(doc)
        doc.add_paragraph().add_run("Бланки согласия для тех, кто ещё не ответил").bold = True
        for r in pack.waiting:
            doc.add_paragraph().add_run(
                "Согласие законного представителя на участие обучающегося в выездном мероприятии").bold = True
            doc.add_paragraph(
                f"Я, ______________________________ (ФИО), законный представитель обучающегося {r.name}, "
                f"{pack.class_title} класс, даю согласие на его (её) участие в выездном мероприятии {pack.title} "
                f"{pack.date_short} ({pack.place}). Сбор в {pack.gathering} {pack.meeting_point}, возвращение к "
                f"{pack.returning}.")
            doc.add_paragraph("Телефон: ______________  Подпись: ____________  Дата: ____________\n")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
